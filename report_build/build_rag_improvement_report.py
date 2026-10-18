from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "IEEE_SMCS_RAG_Improvement_Report.docx"

BLUE = "185ABD"
DARK = "14213D"
MUTED = "657187"
LIGHT = "EDF3FB"
PALE = "F5F7FA"
GREEN = "1B7A55"
GOLD = "9A6B14"
WHITE = "FFFFFF"


def font(run, size=None, bold=None, color=DARK, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def shade(element, fill):
    props = element.get_or_add_tcPr() if element.tag.endswith("tc") else element.get_or_add_pPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    tbl_pr.append(indent)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[i]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell._tc, BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(value), 9, True, WHITE)
    for r_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            if r_index % 2:
                shade(cells[i]._tc, PALE)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            font(p.add_run(str(value)), 9, i == 0)
            if i > 0 and len(str(value)) < 18:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.1
    font(p.add_run(text), 10.5)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    font(p.add_run(text), 10.5)
    return p


def add_callout(doc, label, text, color=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.line_spacing = 1.1
    shade(p._p, LIGHT)
    font(p.add_run(label.upper() + "  "), 9, True, color)
    font(p.add_run(text), 10.5, False, DARK)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    if bold_lead and text.startswith(bold_lead):
        font(p.add_run(bold_lead), 10.5, True)
        font(p.add_run(text[len(bold_lead):]), 10.5)
    else:
        font(p.add_run(text), 10.5)
    return p


def configure(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, before, after, color in (
        ("Heading 1", 16, 16, 8, BLUE), ("Heading 2", 13, 12, 6, BLUE),
        ("Heading 3", 11.5, 8, 4, DARK)):
        style = styles[name]
        style.font.name, style.font.size, style.font.bold = "Calibri", Pt(size), True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)
        style.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(header.add_run("IEEE SMCS  |  RAG SYSTEM IMPROVEMENT"), 8, True, MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("Internal team report  •  3 September 2026"), 8, False, MUTED)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(54)
    p.paragraph_format.space_after = Pt(8)
    font(p.add_run("TECHNICAL IMPROVEMENT REPORT"), 9, True, BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    font(p.add_run("IEEE SMCS Authority-Aware\nRAG Assistant"), 28, True, DARK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(24)
    font(p.add_run("Frontend redesign, production backend hardening, caching, and comparative retrieval research"), 14, False, MUTED)
    add_table(doc, ["Prepared for", "Report date", "System version"],
              [["COMP6015 project team", "3 September 2026", "Backend API 0.3.0"]], [3600, 2880, 2880])
    add_callout(doc, "Executive takeaway", "The project has progressed from a functional RAG prototype into a polished, measurable and safer authority-aware governance assistant. The latest retrieval study shows that more complex retrieval does not automatically outperform the tuned baseline; BM25 and RRF retained Recall@5 while reducing latency, creating a defensible foundation for the next cross-encoder experiment.")
    doc.add_heading("Purpose and scope", level=1)
    add_body(doc, "This report records the complete improvement process completed during the current development cycle. It is intended to help team members understand what changed, why each decision was made, how the system was validated, and which research work should follow.")
    add_bullet(doc, "Product experience: a complete responsive redesign of the Next.js interface.")
    add_bullet(doc, "Server engineering: reusable startup services, health checks, safer production controls, and observability.")
    add_bullet(doc, "Performance: authority-safe response caching with measurable cache behavior.")
    add_bullet(doc, "RAG research: isolated BM25, dense, and reciprocal-rank-fusion retrieval experiments.")

    doc.add_page_break()
    doc.add_heading("1. Starting position", level=1)
    add_body(doc, "The initial system already contained important governance-specific safeguards. Retrieval admitted only approved, production, retrieval-enabled documents; answers carried structured citations; insufficient evidence caused abstention; and explicit document authority governed cross-document conflicts.")
    add_table(doc, ["Existing capability", "Value", "Observed limitation"], [
        ["Deterministic retrieval", "Reproducible and offline", "Relied on hand-tuned lexical rules"],
        ["Hybrid mode", "Added semantic discovery", "Did not isolate sparse and dense contributions"],
        ["Citation validation", "Rejected fabricated evidence", "Validated evidence presence, not every generated claim"],
        ["Golden Set", "Supported repeatable development checks", "Small and repeatedly inspected; not a final unseen test"],
        ["FastAPI endpoints", "Clear chat, audit and evaluation workflows", "Services were reconstructed on each request"],
    ], [2200, 2800, 4360])
    doc.add_heading("System safety boundary", level=2)
    add_body(doc, "Only records satisfying all three conditions enter production retrieval: usage = production, approval_status = approved, and retrieval_enabled = true. Pending documents, audit inputs and evaluation material remain outside the searchable corpus.")
    add_callout(doc, "Design principle", "Authority is applied only after relevance has identified a genuinely related passage. A high-authority but irrelevant document must not win merely because of its authority score.")

    doc.add_heading("2. Frontend product redesign", level=1)
    add_body(doc, "The interface was restructured from a flat prototype into a governance workspace. The redesign preserved all API behavior while making trust, evidence and navigation easier to understand.")
    add_table(doc, ["Area", "Improvement"], [
        ["Navigation", "Desktop sidebar with responsive mobile bottom navigation"],
        ["Primary workflow", "Assistant-first landing experience with sample governance questions"],
        ["Evidence", "Scannable citation cards showing section, page, version and authority"],
        ["Safety states", "Distinct insufficient-evidence, fallback, conflict and human-review treatments"],
        ["Secondary workflows", "Consistent Policy Audit, Document Library and Evaluation views"],
        ["Accessibility", "Semantic controls, visible focus behavior and reduced-motion support"],
    ], [2200, 7160])

    doc.add_page_break()
    doc.add_heading("3. Backend production improvements", level=1)
    doc.add_heading("3.1 Reusable application service graph", level=2)
    add_body(doc, "FastAPI now initializes the registry, search index, embeddings, RAG service, audit service and evaluation service once through the application lifespan. Subsequent requests reuse these objects rather than rebuilding costly retrieval resources.")
    add_body(doc, "Impact: lower repeated-request latency, less memory churn, predictable model initialization, and a clearer dependency graph for testing.")
    doc.add_heading("3.2 Health and readiness", level=2)
    add_table(doc, ["Endpoint", "Purpose", "Key output"], [
        ["/api/health/live", "Confirms the process is responsive", "status"],
        ["/api/health/ready", "Confirms corpus and retrieval services are usable", "documents, chunks, embedding, cache entries"],
        ["/api/health", "Backward-compatible frontend check", "mode and embedding configuration"],
    ], [2300, 4000, 3060])
    doc.add_heading("3.3 Production controls and diagnostics", level=2)
    add_bullet(doc, "CORS origins are configurable through CORS_ORIGINS.")
    add_bullet(doc, "Evaluation endpoints default to disabled when APP_ENV=production.")
    add_bullet(doc, "Every response includes an X-Request-ID for end-to-end tracing.")
    add_bullet(doc, "Structured server logs record method, path, status and latency without logging full policy content.")
    add_bullet(doc, "The API exposes Server-Timing so browser developer tools can separate cache and RAG work.")

    doc.add_heading("4. Authority-safe response cache", level=1)
    add_body(doc, "A bounded, thread-safe TTL/LRU cache was added for chat responses. Its key combines the normalized question with the selected retrieval mode, preventing one experimental mode from receiving another mode’s answer.")
    add_table(doc, ["Property", "Implementation"], [
        ["Default lifetime", "300 seconds"], ["Default capacity", "256 answers"],
        ["Eviction", "Least recently used"], ["Concurrency", "Protected by a re-entrant lock"],
        ["Mutation safety", "Deep copies on cache read and write"],
        ["Invalidation", "Automatic on expiry and process restart/re-index deployment"],
        ["Observability", "X-Cache: HIT/MISS and Server-Timing headers"],
    ], [2500, 6860])
    add_callout(doc, "Why not realtime streaming?", "Deterministic and hybrid responses are already short and citation-gated. Streaming unvalidated LLM tokens would weaken the governance safety model. Streaming should be considered later only as validated status stages, not raw answer tokens.", GOLD)

    doc.add_heading("5. Comparative retrieval experiment", level=1)
    add_body(doc, "The research milestone introduced four directly comparable retrieval configurations while keeping chunking, corpus eligibility, authority handling and downstream audit reasoning unchanged. This ablation design isolates the effect of retrieval strategy.")
    add_table(doc, ["Mode", "Method", "Research role"], [
        ["Deterministic", "Existing lexical coverage and reranking", "Verified project baseline"],
        ["BM25", "Probabilistic term-frequency ranking", "Sparse retrieval ablation"],
        ["Dense", "Local LSA embedding similarity only", "Semantic retrieval ablation"],
        ["RRF", "Reciprocal Rank Fusion of BM25 and dense ranks", "Rank-fusion experiment"],
    ], [1800, 3900, 3660])
    doc.add_heading("5.1 Implementation details", level=2)
    add_bullet(doc, "BM25 uses k1 = 1.5 and b = 0.75 and is implemented locally without a new dependency.")
    add_bullet(doc, "Dense mode uses the configured embedding backend without lexical score boosts.")
    add_bullet(doc, "RRF combines independent BM25 and dense ranks using rank constant k = 60.")
    add_bullet(doc, "Zero-signal results do not receive artificial RRF relevance merely because they occupy a rank position.")
    add_bullet(doc, "Experimental modes exclude the existing project-specific hybrid boosts, keeping the comparison clean.")
    doc.add_heading("5.2 Evaluation measures", level=2)
    add_body(doc, "Retrieval relevance is joined with evaluation-only corrected evidence after prediction, preserving the leakage boundary. Metrics now include:")
    add_bullet(doc, "Recall@5 — proportion of evaluable cases with a relevant passage in the top five.")
    add_bullet(doc, "Mean Reciprocal Rank — rewards placing the first relevant result earlier.")
    add_bullet(doc, "nDCG@10 — applies logarithmic discount to relevant results appearing lower in the ranking.")
    add_bullet(doc, "Detection recall and latency — verify downstream behavior and operational cost.")

    doc.add_heading("6. Initial results", level=1)
    add_table(doc, ["Mode", "Recall@5", "MRR", "nDCG@10", "Detection", "Latency"], [
        ["Deterministic", "90%", "0.833", "0.850", "100%", "473.8 ms"],
        ["BM25", "90%", "0.825", "0.843", "100%", "206.6 ms"],
        ["Dense", "80%", "0.800", "0.800", "100%", "169.7 ms"],
        ["RRF", "90%", "0.825", "0.843", "100%", "193.0 ms"],
    ], [1800, 1450, 1450, 1600, 1600, 1460])
    add_callout(doc, "Finding", "The existing deterministic baseline produced the best ranking metrics in this small development run. BM25 and RRF retained Recall@5 with substantially lower latency. Dense-only retrieval missed one additional evaluable case. Therefore, added semantic complexity did not automatically improve this corpus.")

    doc.add_heading("7. Interpretation and research value", level=1)
    add_body(doc, "The absence of an immediate RRF accuracy gain is a useful research outcome, not a failed implementation. Governance corpora contain exact role names, normative wording, acronyms and quantities that favor sparse lexical matching. The current dense backend is local LSA rather than a domain-trained transformer, so its lower recall is plausible.")
    add_body(doc, "The results support a measured next step: retrieve broadly with BM25 or RRF, then apply a stronger relevance model only to a small candidate set. This preserves recall while testing whether semantic reranking improves the order of genuinely relevant clauses.")
    doc.add_heading("Threats to validity", level=2)
    add_bullet(doc, "The 12-case Golden Set is a development benchmark and has been repeatedly inspected.")
    add_bullet(doc, "The reported latency values are one local run and should not be treated as stable production benchmarks.")
    add_bullet(doc, "A single corrected passage is used as the relevance reference; other passages may also be legitimately relevant.")
    add_bullet(doc, "Local LSA is reproducible and private but weaker than modern domain-tuned dense encoders.")
    add_bullet(doc, "Authority and approval metadata still require governance-owner verification before production use.")
    add_bullet(doc, "Downstream deterministic classification can mask retrieval differences when several candidates support the same outcome.")

    doc.add_heading("8. Verification completed", level=1)
    add_table(doc, ["Check", "Result"], [
        ["Backend unit and integration tests", "60 passed"],
        ["Frontend workflow tests", "5 passed"],
        ["Frontend lint", "Passed"],
        ["Next.js production build", "Passed"],
        ["Python compilation", "Passed"],
        ["Git whitespace validation", "Passed"],
    ], [4200, 5160])

    doc.add_heading("9. Recommended next milestone", level=1)
    add_callout(doc, "Recommendation", "Add a cross-encoder reranker over the top 20 BM25/RRF candidates and evaluate whether it improves MRR and nDCG@10 without reducing Recall@5.")
    add_number(doc, "Freeze the present four-mode results and configuration as the baseline.")
    add_number(doc, "Retrieve the top 20 candidates using BM25 and RRF.")
    add_number(doc, "Rerank candidates with a general or legal-domain cross-encoder.")
    add_number(doc, "Measure Recall@5, MRR, nDCG@10, abstention, latency and downstream answer correctness.")
    add_number(doc, "Run ablations with authority applied only after relevance, then as a small reranking feature.")
    add_number(doc, "Validate the selected configuration on the untouched 50-case holdout only after tuning is complete.")

    doc.add_heading("10. Suggested team responsibilities", level=1)
    add_table(doc, ["Workstream", "Suggested owner activity", "Deliverable"], [
        ["Retrieval", "Implement and profile cross-encoder reranking", "Reproducible experiment branch"],
        ["Evaluation", "Review relevance judgments and protect holdout isolation", "Signed evaluation protocol"],
        ["Governance", "Verify authority, dates, approval and supersession metadata", "Approved registry revision"],
        ["Frontend", "Add per-case rank and retrieval diagnostics", "Research dashboard update"],
        ["Documentation", "Record configurations, seeds and hardware", "Experiment log and final report"],
    ], [1800, 4900, 2660])

    doc.add_heading("Appendix A — Configuration and endpoints", level=1)
    doc.add_heading("Environment controls", level=2)
    add_table(doc, ["Variable", "Default", "Purpose"], [
        ["APP_ENV", "development", "Controls development-only behavior"],
        ["CORS_ORIGINS", "http://localhost:3000", "Permitted frontend origins"],
        ["ENABLE_EVALUATION_API", "true in development", "Protects expensive evaluation endpoints"],
        ["CHAT_CACHE_TTL_SECONDS", "300", "Cached-answer lifetime"],
        ["CHAT_CACHE_MAX_ENTRIES", "256", "Maximum cached answers"],
        ["EMBEDDING_BACKEND", "local_lsa", "Dense retrieval implementation"],
        ["MIN_HYBRID_SCORE", "0.12", "Non-deterministic relevance threshold"],
    ], [3100, 2400, 3860])
    doc.add_heading("Key endpoints", level=2)
    for item in ("GET /api/health", "GET /api/health/live", "GET /api/health/ready",
                 "POST /api/chat", "POST /api/audit", "POST /api/evaluation/compare",
                 "POST /api/debug/retrieval — development only"):
        add_bullet(doc, item)
    doc.add_heading("Appendix B — Reproduction commands", level=1)
    for command in (
        "./venv/bin/python -m uvicorn backend.main:app --reload --port 8000",
        "npm run dev --prefix frontend",
        "./venv/bin/python -m unittest tests.test_response_cache tests.test_vertical_slice tests.test_authority_conflicts tests.test_hybrid_rag -v",
        "npm test --prefix frontend && npm run lint --prefix frontend && npm run build --prefix frontend",
    ):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(6)
        shade(p._p, PALE)
        run = p.add_run(command)
        run.font.name = "Courier New"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Courier New")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(DARK)

    doc.core_properties.title = "IEEE SMCS Authority-Aware RAG Assistant — Technical Improvement Report"
    doc.core_properties.subject = "Frontend, backend, caching and retrieval research improvements"
    doc.core_properties.author = "COMP6015 Project Team"
    doc.core_properties.keywords = "RAG, IEEE SMCS, governance, BM25, dense retrieval, reciprocal rank fusion"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
